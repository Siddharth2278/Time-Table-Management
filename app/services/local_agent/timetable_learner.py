"""Deterministic timetable learner (no model calls in Phase 1).

Reads historical timetables from CSV / XLSX / JSON and extracts
college-specific scheduling patterns keyed by structural ROLES
(type class + duration + frequency) — never by subject name — so the
learned profile transfers to new subjects in Phase 2.
"""
import csv
import json
from collections import Counter
from typing import Any, Dict, List, Tuple

from app.services.local_agent.schemas import (
    LearningError, LectureRow, blank_profile,
)
from app.utils.helpers import time_to_minutes

MIN_LECTURES = 3

FIELD_ALIASES = {
    "code": {"code", "subject_code", "subjectcode", "sub_code", "course_code"},
    "name": {"name", "subject_name", "subjectname", "course", "course_name"},
    "type": {"type", "subject_type", "subjecttype", "lecture_type"},
    "duration": {"duration", "lecture_duration", "mins", "minutes"},
    "day": {"day", "day_name", "weekday"},
    "start": {"start", "start_time", "from", "from_time"},
    "end": {"end", "end_time", "to", "to_time"},
    "teacher": {"teacher", "teacher_name", "faculty", "staff"},
    "room": {"room", "room_name", "room_no", "room_number", "lab"},
    "break": {"break", "is_break", "break_name", "lunch"},
}


def _normalize_row(raw: dict) -> LectureRow:
    lowered = {(str(k) or "").strip().lower(): v for k, v in dict(raw).items()}
    found: Dict[str, str] = {}
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            if alias in lowered and lowered[alias] not in (None, ""):
                found[field] = str(lowered[alias]).strip()
                break
        else:
            found[field] = ""
    is_break = found["break"].strip().lower() in (
        "1", "true", "yes", "y", "break", "lunch")
    try:
        duration = int(float(found["duration"] or 60))
    except (TypeError, ValueError):
        duration = 60
    row = LectureRow(
        code=found["code"], name=found["name"],
        type=(found["type"] or "Theory").strip() or "Theory",
        duration=duration if duration > 0 else 60,
        day=found["day"], start=found["start"], end=found["end"],
        teacher=found["teacher"], room=found["room"])
    row_break = is_break
    row.__dict__["is_break"] = row_break
    return row


def _read_json(path: str) -> List[dict]:
    try:
        with open(path, "r", encoding="utf-8-sig") as fh:
            data = json.load(fh)
    except (OSError, ValueError) as e:
        raise LearningError(f"Cannot read timetable file '{path}': {e}")
    if isinstance(data, dict) and isinstance(data.get("entries"), list):
        data = data["entries"]
    if isinstance(data, dict) and isinstance(data.get("lectures"), list):
        data = data["lectures"]
    if not isinstance(data, list) or not all(isinstance(r, dict) for r in data):
        raise LearningError(
            f"Invalid timetable file '{path}': JSON must be a list of lecture rows.")
    return data


def _read_csv(path: str) -> List[dict]:
    try:
        with open(path, "r", encoding="utf-8-sig", newline="") as fh:
            return list(csv.DictReader(fh))
    except OSError as e:
        raise LearningError(f"Cannot read timetable file '{path}': {e}")


def _read_xlsx(path: str) -> List[dict]:
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows: List[dict] = []
        headers: List[str] = []
        for values in sheet.iter_rows(values_only=True):
            if not headers:
                headers = [(str(h) or "").strip() for h in values]
                if not any(headers):
                    raise LearningError(
                        f"Invalid timetable file '{path}': first row must hold headers.")
                continue
            if all(v is None or str(v).strip() == "" for v in values):
                continue
            rows.append({h: ("" if v is None else str(v)) for h, v in zip(headers, values)})
        workbook.close()
        return rows
    except LearningError:
        raise
    except ImportError:
        raise LearningError("Excel support needs the openpyxl package.")
    except Exception as e:
        raise LearningError(f"Cannot read timetable file '{path}': {e}")


def load_rows(file_path: str) -> Tuple[List[LectureRow], int]:
    """Load + normalize rows. Returns (valid_rows, skipped_count)."""
    suffix = str(file_path).lower().rsplit(".", 1)[-1] if "." in str(file_path) else ""
    if suffix == "json":
        raw = _read_json(str(file_path))
    elif suffix == "csv":
        raw = _read_csv(str(file_path))
    elif suffix in ("xlsx", "xlsm", "xltx", "xltm"):
        raw = _read_xlsx(str(file_path))
    else:
        raise LearningError(
            f"Unsupported timetable file '{file_path}' (use .csv, .json or .xlsx).")
    rows: List[LectureRow] = []
    skipped = 0
    for raw_row in raw:
        try:
            row = _normalize_row(raw_row)
        except Exception:
            skipped += 1
            continue
        if row.__dict__.get("is_break"):
            skipped += 1
            continue
        try:
            start, end = time_to_minutes(row.start), time_to_minutes(row.end)
        except (ValueError, AttributeError, TypeError):
            skipped += 1
            continue
        if end <= start or not row.code.strip() or not row.day.strip():
            skipped += 1
            continue
        rows.append(row)
    return rows, skipped


def _role_key(row: LectureRow) -> str:
    kind = "practical" if (row.type or "Theory") != "Theory" else "theory"
    return f"{kind}|{row.duration}"


def _ranked(counter: Counter) -> List[str]:
    return [k for k, _ in sorted(counter.items(), key=lambda kv: (-kv[1], kv[0]))]


def learn(rows: List[LectureRow], source_label: str = "") -> Dict[str, Any]:
    """Extract the structural learning profile (deterministic, offline)."""
    if len(rows) < MIN_LECTURES:
        raise LearningError(
            f"Insufficient timetable data: found {len(rows)} valid lecture "
            f"row(s), need at least {MIN_LECTURES} to learn patterns.")
    by_code: Dict[str, List[LectureRow]] = {}
    for row in rows:
        by_code.setdefault(row.code.strip(), []).append(row)

    subjects = []
    role_stats: Dict[str, Dict[str, Any]] = {}
    for code in sorted(by_code):
        group = by_code[code]
        first = group[0]
        kind = "practical" if (first.type or "Theory") != "Theory" else "theory"
        days = sorted({r.day.strip() for r in group})
        times = sorted({r.start.strip() for r in group})
        morning = afternoon = 0
        for r in group:
            try:
                if time_to_minutes(r.start) < 12 * 60:
                    morning += 1
                else:
                    afternoon += 1
            except (ValueError, AttributeError, TypeError):
                continue
        subjects.append({
            "code": code, "name": first.name, "type": first.type,
            "duration": first.duration, "count": len(group),
            "days": days, "times": times,
            "morning": morning, "afternoon": afternoon,
        })
        key = f"{kind}|{first.duration}|{len(group)}"
        role = role_stats.setdefault(key, {
            "type": kind, "duration": first.duration, "frequency": len(group),
            "subjects": 0, "day_counts": [], "day_votes": Counter(),
            "time_votes": Counter(), "morning": 0, "afternoon": 0,
            "adjacent_pairs": 0, "pairs_seen": 0,
        })
        role["subjects"] += 1
        role["day_counts"].append(len(days))
        for day in days:
            role["day_votes"][day] += 1
        for time in times:
            role["time_votes"][time] += 1
        role["morning"] += morning
        role["afternoon"] += afternoon
        per_day: Dict[str, list] = {}
        for r in group:
            try:
                per_day.setdefault(r.day.strip(), []).append(time_to_minutes(r.start))
            except (ValueError, AttributeError, TypeError):
                continue
        for starts in per_day.values():
            ordered = sorted(starts)
            for a, b in zip(ordered, ordered[1:]):
                role["pairs_seen"] += 1
                if 0 < b - a <= 70:
                    role["adjacent_pairs"] += 1

    day_load: Counter = Counter()
    time_load: Counter = Counter()
    teacher_load: Counter = Counter()
    room_load: Counter = Counter()
    morning_total = afternoon_total = 0
    for row in rows:
        day_load[row.day.strip()] += 1
        time_load[row.start.strip()] += 1
        teacher = row.teacher.strip() or "Unknown"
        teacher_load[teacher] += 1
        room = row.room.strip() or "Unknown"
        room_load[room] += 1
        try:
            if time_to_minutes(row.start) < 12 * 60:
                morning_total += 1
            else:
                afternoon_total += 1
        except (ValueError, AttributeError, TypeError):
            continue
    total = len(rows)
    loads = sorted(teacher_load.values()) or [0]

    roles = {}
    for key in sorted(role_stats):
        stat = role_stats[key]
        day_counts = stat["day_counts"] or [0]
        roles[key] = {
            "type": stat["type"],
            "duration": stat["duration"],
            "frequency": stat["frequency"],
            "subjects_observed": stat["subjects"],
            "avg_day_count": round(sum(day_counts) / len(day_counts), 2),
            "preferred_days": _ranked(stat["day_votes"]),
            "preferred_times": _ranked(stat["time_votes"]),
            "morning_share": round(stat["morning"] / max(1, stat["morning"] + stat["afternoon"]), 3),
            "adjacent_pair_rate": round(
                stat["adjacent_pairs"] / max(1, stat["pairs_seen"]), 3),
        }
    teacher_counts = sorted(teacher_load.values())
    room_counts = sorted(room_load.values())
    profile = blank_profile(source_label)
    profile.update({
        "total_lectures": total,
        "subjects": subjects,
        "roles": roles,
        "patterns": {
            "preferred_days": _ranked(day_load),
            "preferred_times": _ranked(time_load),
            "morning_share": round(morning_total / max(1, total), 3),
            "average_daily_load": round(total / max(1, len(day_load)), 2),
            "min_daily_load": min(day_load.values()) if day_load else 0,
            "max_daily_load": max(day_load.values()) if day_load else 0,
            "daily_load": dict(sorted(day_load.items())),
            "time_load": dict(sorted(time_load.items())),
            "teacher_workload": dict(sorted(teacher_load.items())),
            "teacher_min": min(teacher_counts) if teacher_counts else 0,
            "teacher_max": max(teacher_counts) if teacher_counts else 0,
            "teacher_avg": round(sum(teacher_counts) / max(1, len(teacher_counts)), 2),
            "room_usage": dict(sorted(room_load.items())),
            "room_min": min(room_counts) if room_counts else 0,
            "room_max": max(room_counts) if room_counts else 0,
            "theory_duration": _typical([r.duration for r in rows if (r.type or "Theory") == "Theory"]),
            "practical_duration": _typical([r.duration for r in rows if (r.type or "Theory") != "Theory"]),
        },
    })
    return profile


def _typical(values: List[int]):
    if not values:
        return None
    return Counter(values).most_common(1)[0][0]
