"""Reference timetable analysis: real counts and distributions, read-only."""
from typing import Any, Dict

from app.models import Semester, Subject, TimeSlot, TimetableEntry, WorkingDay
from app.utils.helpers import time_to_minutes


def _shape_profile(entries, semester_label, working_days, slots, breaks):
    """Shared shaping for DB rows and external file rows.

    Each entry: {code, name, type, duration, day, start, end, teacher, room}.
    """
    subjects = {}
    day_dist = {}
    time_dist = {}
    teacher_load = {}
    room_load = {}
    practical = 0
    by_day_subject = {}
    morning = afternoon = 0
    skipped = 0
    for e in entries:
        try:
            mins = time_to_minutes(e["start"])
            emins = time_to_minutes(e["end"])
            if emins <= mins:
                raise ValueError("bad range")
        except (ValueError, AttributeError, TypeError, KeyError):
            skipped += 1
            continue
        code = (e.get("code") or "").strip()
        if not code or not (e.get("day") or "").strip():
            skipped += 1
            continue
        ltype = (e.get("type") or "Theory").strip() or "Theory"
        try:
            duration = int(e.get("duration") or 60)
        except (TypeError, ValueError):
            duration = 60
        info = subjects.setdefault(code, {
            "code": code, "name": (e.get("name") or ""),
            "type": ltype, "duration": duration, "count": 0,
        })
        info["count"] += 1
        day = e["day"].strip()
        day_dist[day] = day_dist.get(day, 0) + 1
        time_dist[e["start"]] = time_dist.get(e["start"], 0) + 1
        teacher = (e.get("teacher") or "").strip() or "Unknown"
        teacher_load[teacher] = teacher_load.get(teacher, 0) + 1
        room = (e.get("room") or "").strip() or "Unknown"
        room_load[room] = room_load.get(room, 0) + 1
        if ltype != "Theory":
            practical += 1
            by_day_subject.setdefault(day, {}).setdefault(code, []).append(e["start"])
        if mins < 12 * 60:
            morning += 1
        else:
            afternoon += 1
    consecutive_practicals = 0
    for per_subject in by_day_subject.values():
        for times in per_subject.values():
            if len(times) < 2:
                continue
            try:
                ordered = sorted(time_to_minutes(t) for t in times)
            except (ValueError, AttributeError, TypeError):
                continue
            for a, b in zip(ordered, ordered[1:]):
                if 0 < b - a <= 70:
                    consecutive_practicals += 1
    total = sum(day_dist.values())
    cells = max(1, len(working_days) * max(1, len(slots)))
    loads = sorted(day_dist.values()) if day_dist else [0]
    return {
        "semester": {"id": None, "name": semester_label},
        "total_lectures": total,
        "subjects": sorted(subjects.values(), key=lambda s: (-s["count"], s["code"])),
        "subject_count": len(subjects),
        "practical_sessions": practical,
        "consecutive_practical_pairs": consecutive_practicals,
        "day_distribution": day_dist,
        "time_distribution": time_dist,
        "morning_lectures": morning,
        "afternoon_lectures": afternoon,
        "working_days": list(working_days),
        "breaks": list(breaks),
        "slots_per_day": len(slots),
        "average_daily_lectures": round(total / max(1, len(day_dist)), 2) if day_dist else 0.0,
        "min_daily_load": min(loads),
        "max_daily_load": max(loads),
        "free_periods": max(0, cells - total),
        "density": round(total / cells, 3),
        "teacher_load": teacher_load,
        "room_load": room_load,
        "first_last": {},
        "has_data": total > 0,
        "skipped_rows": skipped,
    }


def _db_collections(session):
    days = session.query(WorkingDay).filter(
        WorkingDay.is_enabled == True).order_by(WorkingDay.sort_order).all()  # noqa: E712
    slots = session.query(TimeSlot).filter(
        TimeSlot.is_enabled == True, TimeSlot.is_break == False  # noqa: E712
    ).order_by(TimeSlot.start_time).all()
    breaks = session.query(TimeSlot).filter(
        TimeSlot.is_enabled == True, TimeSlot.is_break == True  # noqa: E712
    ).order_by(TimeSlot.start_time).all()
    return ([d.name for d in days], len(slots),
            [{"start": b.start_time, "end": b.end_time,
              "name": b.break_name or "Break"} for b in breaks])


def _day_name(session, day_id):
    day = session.query(WorkingDay).filter(WorkingDay.id == day_id).first()
    return day.name if day else f"Day {day_id}"


def analyze_reference(session, semester_id: int) -> Dict[str, Any]:
    """Structured ReferenceProfile of an existing timetable (possibly empty)."""
    sem = session.query(Semester).filter(Semester.id == semester_id).first()
    entries = session.query(TimetableEntry).filter(
        TimetableEntry.semester_id == semester_id).all()
    rows = []
    for e in entries:
        sub = session.query(Subject).filter(Subject.id == e.subject_id).first()
        rows.append({
            "code": sub.code if sub else f"Sub {e.subject_id}",
            "name": sub.name if sub else "",
            "type": sub.subject_type if sub else e.lecture_type,
            "duration": sub.lecture_duration if sub else 60,
            "day": _day_name(session, e.day_id),
            "start": e.start_time, "end": e.end_time,
            "teacher": e.teacher.name if e.teacher else f"Teacher {e.teacher_id}",
            "room": e.room.name if e.room else f"Room {e.room_id}",
        })
    working_days, slot_count, breaks = _db_collections(session)
    profile = _shape_profile(rows, sem.name if sem else f"Semester {semester_id}",
                             working_days, ["x"] * slot_count, breaks)
    profile["semester"] = {"id": semester_id, "name": sem.name if sem else ""}
    # Per-day first/last lecture (DB-only detail).
    first_last: Dict[str, Dict[str, Any]] = {}
    for e in entries:
        try:
            mins = time_to_minutes(e.start_time)
            emins = time_to_minutes(e.end_time)
        except (ValueError, AttributeError, TypeError):
            continue
        day_name = _day_name(session, e.day_id)
        slot = first_last.setdefault(day_name, {"first": e.start_time, "first_mins": mins,
                                                 "last": e.end_time, "last_mins": emins})
        if mins < slot["first_mins"]:
            slot["first"], slot["first_mins"] = e.start_time, mins
        if emins > slot["last_mins"]:
            slot["last"], slot["last_mins"] = e.end_time, emins
    profile["first_last"] = {d: {"first": v["first"], "last": v["last"]}
                             for d, v in first_last.items()}
    return profile


FIELD_ALIASES = {
    "code": {"code", "subject_code", "subjectcode", "sub_code", "course_code"},
    "name": {"name", "subject_name", "subjectname", "course", "course_name"},
    "type": {"type", "subject_type", "subjecttype", "lecture_type"},
    "duration": {"duration", "lecture_duration", "mins", "minutes"},
    "day": {"day", "day_name", "weekday"},
    "start": {"start", "start_time", "from", "from_time"},
    "end": {"end", "end_time", "to", "to_time"},
    "teacher": {"teacher", "teacher_name", "faculty", "staff"},
    "room": {"room", "room_name", "room_no", "room_number", "lab"},
}


def _normalize_row(raw: dict) -> dict:
    lowered = {(str(k) or "").strip().lower(): v for k, v in dict(raw).items()}
    row = {}
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            if alias in lowered and lowered[alias] not in (None, ""):
                row[field] = str(lowered[alias]).strip()
                break
        else:
            row[field] = ""
    return row


def load_reference_file(path: str):
    """Load an EXTERNAL reference timetable (csv / json / xlsx).

    Returns (rows, skipped). Raises IntelligenceError when nothing usable.
    Column/keys are matched case-insensitively with common aliases.
    """
    from app.services.intelligence.timetable_agent import IntelligenceError
    suffix = str(path).lower().rsplit(".", 1)[-1] if "." in str(path) else ""
    raw_rows: list = []
    try:
        if suffix == "json":
            import json
            with open(path, "r", encoding="utf-8-sig") as fh:
                data = json.load(fh)
            if isinstance(data, dict) and isinstance(data.get("entries"), list):
                data = data["entries"]
            if not isinstance(data, list):
                raise ValueError("JSON must be a list of rows.")
            raw_rows = [r for r in data if isinstance(r, dict)]
        elif suffix == "csv":
            import csv
            with open(path, "r", encoding="utf-8-sig", newline="") as fh:
                raw_rows = list(csv.DictReader(fh))
        elif suffix in ("xlsx", "xlsm", "xltx", "xltm"):
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook.active
            headers = []
            for row in sheet.iter_rows(values_only=True):
                if not headers:
                    headers = [(str(h) or "").strip() for h in row]
                    continue
                if all(v is None or str(v).strip() == "" for v in row):
                    raw_rows.append({})
                    continue
                raw_rows.append({h: ("" if v is None else str(v)) for h, v in zip(headers, row)})
            workbook.close()
        else:
            raise ValueError("Unsupported file type (use .csv, .json or .xlsx).")
    except IntelligenceError:
        raise
    except Exception as e:
        raise IntelligenceError(f"Could not read reference file: {e}")
    rows = [_normalize_row(r) for r in raw_rows]
    valid = [r for r in rows if r["code"] and r["day"] and r["start"] and r["end"]]
    if not valid:
        raise IntelligenceError(
            "No usable lecture rows found. Need columns like: "
            "subject_code, day, start_time, end_time.")
    return valid, len(rows) - len(valid)


def analyze_external_reference(session, rows, label: str) -> Dict[str, Any]:
    """Profile external rows against current working days/slots/breaks."""
    working_days, slot_count, breaks = _db_collections(session)
    profile = _shape_profile(rows, label, working_days, ["x"] * slot_count, breaks)
    return profile
