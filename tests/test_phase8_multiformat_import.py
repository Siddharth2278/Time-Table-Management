"""Phase 8 tests: multi-format import + offline extraction.

Small deterministic fixtures, no internet. OCR execution tests run only
when a local Tesseract runtime exists; the importer interface,
preprocessing, normalization, validation and error paths are always tested
(real coverage without the binary via monkeypatched word lists).
"""
import csv
import io
import os
import zipfile

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from pathlib import Path

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.models import (
    Base, Room, Semester, Subject, Teacher, TimeSlot, TimetableEntry, WorkingDay,
)
from app.services.conflict_service import ConflictService
from app.services.local_agent.agent import TimetableAgent
from app.services.local_agent.model_client import OllamaClient
from app.services.local_agent.schemas import LearningError
from app.services.timetable_import.exceptions import TimetableImportError
from app.services.timetable_import.importer import TimetableImporter
from app.services.timetable_import import table_normalizer as norm

ROWS_A = [
    {"Day": "Monday", "Start": "09:00", "End": "10:00",
     "Subject": "Java", "Teacher": "Dr A", "Room": "101"},
    {"Day": "Wednesday", "Start": "09:00", "End": "10:00",
     "Subject": "Java", "Teacher": "Dr A", "Room": "101"},
    {"Day": "Friday", "Start": "09:00", "End": "10:00",
     "Subject": "Java", "Teacher": "Dr A", "Room": "101"},
    {"Day": "Tuesday", "Start": "14:00", "End": "16:00",
     "Subject": "DBMS Lab", "Teacher": "Dr B", "Room": "Lab 1"},
    {"Day": "Thursday", "Start": "14:00", "End": "16:00",
     "Subject": "DBMS Lab", "Teacher": "Dr B", "Room": "Lab 1"},
]

GRID_B = [
    ["Time", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday"],
    ["09:00-10:00", "Java", "", "Java", "", "Java"],
    ["10:00-11:00", "", "OS", "", "OS", ""],
    ["14:00-16:00", "", "DBMS Lab", "", "DBMS Lab", ""],
]

GRID_C = [
    ["Day", "Time", "Subject"],
    ["Monday", "9-10", "Java"],
    ["Wednesday", "09:00 - 10:00", "Java"],
    ["Friday", "9:00 AM - 10:00 AM", "Java"],
]


def write_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    return str(path)


def write_grid_csv(path, grid):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows(grid)
    return str(path)


def write_xlsx(path, grid, merged=None):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    for row in grid:
        sheet.append(row)
    for coord in (merged or []):
        sheet.merge_cells(coord)
    workbook.save(path)
    workbook.close()
    return str(path)


def make_session():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine)()


def seed(session):
    sem = Semester(name="Semester 1", status="Active")
    session.add(sem)
    session.flush()
    teachers = [Teacher(name=f"T{i}", email=f"t{i}@c.edu", department="CS",
                        status="Active") for i in (1, 2)]
    rooms = [Room(name=f"R{i}", room_number=f"10{i}", type="Classroom",
                  status="Available") for i in (1, 2)]
    session.add_all(teachers + rooms)
    session.flush()
    days = [WorkingDay(name=n, is_enabled=(n != "Sunday"), sort_order=i)
            for i, n in enumerate(["Monday", "Tuesday", "Wednesday",
                                   "Thursday", "Friday", "Saturday",
                                   "Sunday"])]
    session.add_all(days)
    slots = [TimeSlot(start_time="09:00", end_time="10:00", label="09:00-10:00"),
             TimeSlot(start_time="10:00", end_time="11:00", label="10:00-11:00"),
             TimeSlot(start_time="11:00", end_time="12:00", label="11:00-12:00"),
             TimeSlot(start_time="14:00", end_time="15:00", label="14:00-15:00"),
             TimeSlot(start_time="15:00", end_time="16:00", label="15:00-16:00")]
    session.add_all(slots)
    subjects = [Subject(code="NW101", name="Networks", semester_id=sem.id,
                        subject_type="Theory",
                        required_lectures_per_week=2, lecture_duration=60,
                        teacher_id=teachers[0].id, room_id=rooms[0].id)]
    session.add_all(subjects)
    session.commit()
    return sem


def ocr_available():
    from app.services.timetable_import import image_importer
    return bool(image_importer.check_ocr_available().get("available"))


# ---- structured formats ----

def test_csv_layout_a_rows(tmp_path):
    path = write_csv(str(tmp_path / "a.csv"), ROWS_A)
    extraction = TimetableImporter.import_file(path)
    assert extraction.lectures == 5
    assert extraction.method == "csv"
    by_day = {r.day: r.subject for r in extraction.records}
    assert by_day["Monday"] == "Java"
    assert extraction.records[3].start == "14:00"


def test_csv_layout_b_grid(tmp_path):
    path = write_grid_csv(str(tmp_path / "b.csv"), GRID_B)
    extraction = TimetableImporter.import_file(path)
    assert extraction.lectures == 7  # 3 Java + 2 OS + 2 DBMS Lab
    assert {r.day for r in extraction.records} >= {"Monday", "Tuesday"}


def test_csv_layout_c_loose_times(tmp_path):
    path = write_grid_csv(str(tmp_path / "c.csv"), GRID_C)
    extraction = TimetableImporter.import_file(path)
    assert extraction.lectures == 3
    assert all(r.start == "09:00" and r.end == "10:00"
               for r in extraction.records)


def test_xlsx_grid_and_merged_cells(tmp_path):
    pytest.importorskip("openpyxl")
    grid = [row[:] for row in GRID_B]
    path = write_xlsx(str(tmp_path / "g.xlsx"), grid, merged=["B2:B3"])
    extraction = TimetableImporter.import_file(path)
    assert extraction.method == "excel"
    assert extraction.lectures >= 7


def test_xls_legacy(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Timetable")
    for r, row in enumerate(GRID_B):
        for c, value in enumerate(row):
            sheet.write(r, c, value)
    path = str(tmp_path / "legacy.xls")
    workbook.save(path)
    extraction = TimetableImporter.import_file(path)
    assert extraction.file_type == ".xls"
    assert extraction.lectures == 7


# ---- day/time normalization ----

def test_day_and_time_normalization():
    assert norm.parse_day("Mon") == "Monday"
    assert norm.parse_day("MON") == "Monday"
    assert norm.parse_day("tue") == "Tuesday"
    assert norm.parse_day("Funday") == ""
    assert norm.parse_time_range("9-10") == ("09:00", "10:00")
    assert norm.parse_time_range("09:00-10:00") == ("09:00", "10:00")
    assert norm.parse_time_range("9:00 AM - 10:00 AM") == ("09:00", "10:00")
    assert norm.parse_time_range("2 PM - 4 PM") == ("14:00", "16:00")
    assert norm.parse_time_range("banana") is None
    assert norm.parse_time_range("10:00-09:00") is None


def test_orientation_days_vertical():
    grid = [
        ["", "09:00-10:00", "10:00-11:00"],
        ["Monday", "Java", "OS"],
        ["Tuesday", "DBMS", "Java"],
    ]
    records, warnings = norm.grid_to_records(grid, "v.csv")
    assert len(records) == 4
    assert {(r.day, r.start, r.subject) for r in records} == {
        ("Monday", "09:00", "Java"), ("Monday", "10:00", "OS"),
        ("Tuesday", "09:00", "DBMS"), ("Tuesday", "10:00", "Java")}


def test_words_to_records_both_orientations():
    def word(text, x, y, conf=95.0):
        return {"text": text, "x0": x, "y0": y, "x1": x + 40,
                "y1": y + 12, "conf": conf}

    across = [word("MON", 120, 0), word("TUE", 220, 0),
              word("09:00-10:00", 0, 20), word("JAVA", 120, 20),
              word("OS", 220, 20)]
    records, _ = norm.words_to_records(across, "img.png")
    assert {(r.day, r.subject) for r in records} == {
        ("Monday", "JAVA"), ("Tuesday", "OS")}

    down = [word("09:00-10:00", 120, 0), word("10:00-11:00", 220, 0),
            word("Monday", 0, 20), word("JAVA", 120, 20),
            word("OS", 220, 20)]
    records, _ = norm.words_to_records(down, "img.png")
    assert {(r.day, r.start, r.subject) for r in records} == {
        ("Monday", "09:00", "JAVA"), ("Monday", "10:00", "OS")}


# ---- confidence + approval ----

def test_confidence_report_and_low_rejection(tmp_path):
    from app.services.timetable_import.models import TimetableRecord
    from app.services.timetable_import.validator import validate_records
    records = [
        TimetableRecord(day="Monday", start="09:00", end="10:00",
                        subject="Java", confidence=0.95),
        TimetableRecord(day="Tuesday", start="09:00", end="10:00",
                        subject="Java", confidence=0.6),
        TimetableRecord(day="Funday", start="09:00", end="10:00",
                        subject="Java", confidence=0.9),
        TimetableRecord(day="Monday", start="09:00", end="10:00",
                        subject="Java", confidence=0.2),
    ]
    approved, rejected = validate_records(records)
    assert len(approved) == 2 and len(rejected) == 2


def test_approval_flow_stages_jsonl(tmp_path, monkeypatch):
    from app.services.timetable_import.preview import save_approved_jsonl
    from app.services.timetable_import.models import TimetableRecord
    monkeypatch.setenv("APPDATA", str(tmp_path))
    records = [TimetableRecord(day="Monday", start="09:00", end="10:00",
                               subject="Java", confidence=0.9,
                               source_file="g.png")]
    path, approved_count, rejected_count = save_approved_jsonl(records)
    assert approved_count == 1 and rejected_count == 0
    assert Path(path).is_file()
    with pytest.raises(LearningError):
        save_approved_jsonl([TimetableRecord(day="Nope", start="xx",
                                             end="yy", subject="",
                                             confidence=0.1)])


# ---- images (offline OCR) ----

def _grid_image(path, grid):
    from PIL import Image, ImageDraw
    cell_w, cell_h = 220, 60
    image = Image.new("RGB",
                      (cell_w * len(grid[0]), cell_h * len(grid)),
                      "white")
    draw = ImageDraw.Draw(image)
    for r, row in enumerate(grid):
        for c, text in enumerate(row):
            draw.text((c * cell_w + 10, r * cell_h + 18), text, fill="black")
            draw.rectangle([c * cell_w, r * cell_h,
                            (c + 1) * cell_w, (r + 1) * cell_h],
                           outline="gray")
    image.save(path)
    return str(path)


def test_image_preprocessing_offline(tmp_path):
    from app.services.timetable_import import image_importer
    path = _grid_image(str(tmp_path / "grid.png"), GRID_B)
    processed = image_importer.preprocess_image(path)
    assert processed.size[0] >= 220 * len(GRID_B[0])
    assert processed.mode in ("L", "RGB")


def test_image_ocr_path_without_binary(monkeypatch, tmp_path):
    from app.services.timetable_import import image_importer
    path = _grid_image(str(tmp_path / "grid.png"), GRID_B)
    monkeypatch.setattr(image_importer, "check_ocr_available",
                        lambda: {"available": False,
                                 "reason": "tesseract-binary-missing"})
    with pytest.raises(TimetableImportError, match="Local OCR is not available"):
        image_importer.extract(path, "grid.png")


def test_image_grid_via_mocked_ocr(monkeypatch, tmp_path):
    from app.services.timetable_import import image_importer

    def fake_ocr(image):
        words = []
        header = ["", "MON", "TUE"]
        rows = [["09:00-10:00", "JAVA", "OS"]]
        for y, line in enumerate([header] + rows):
            for x, text in enumerate(line):
                if text:
                    words.append({"text": text, "x0": x * 100, "y0": y * 30,
                                  "x1": x * 100 + 60, "y1": y * 30 + 15,
                                  "conf": 96.0})
        return words

    monkeypatch.setattr(image_importer, "ocr_words", fake_ocr)
    path = _grid_image(str(tmp_path / "grid.png"), GRID_B)
    extraction = image_importer.extract(path, "grid.png")
    assert extraction.method == "ocr"
    assert {(r.day, r.subject) for r in extraction.records} == {
        ("Monday", "JAVA"), ("Tuesday", "OS")}


def test_jpeg_import_real_ocr(tmp_path):
    pytest.importorskip("pytesseract")
    from app.services.timetable_import import image_importer
    if not image_importer.check_ocr_available().get("available"):
        pytest.skip("Local Tesseract runtime is not installed.")
    path = _grid_image(str(tmp_path / "grid.jpg"), GRID_B)
    extraction = TimetableImporter.import_file(path)
    assert extraction.lectures >= 1


# ---- PDF ----

def _text_pdf(path):
    from reportlab.pdfgen import canvas
    doc = canvas.Canvas(path)
    doc.setFont("Helvetica", 12)
    y = 750
    for line in ["Time Monday Tuesday",
                 "09:00-10:00 JAVA OS",
                 "10:00-11:00 DBMS JAVA"]:
        doc.drawString(50, y, line)
        y -= 20
    doc.save()
    return str(path)


def test_pdf_text_extraction(tmp_path):
    pytest.importorskip("pymupdf")
    path = _text_pdf(str(tmp_path / "t.pdf"))
    extraction = TimetableImporter.import_file(path)
    assert extraction.file_type == ".pdf"
    assert "pdf-text" in extraction.method
    assert extraction.pages == 1
    assert extraction.lectures >= 2


def test_scanned_pdf_ocr_path(monkeypatch, tmp_path):
    pytest.importorskip("pymupdf")
    from app.services.timetable_import import image_importer
    from reportlab.pdfgen import canvas
    from PIL import Image, ImageDraw
    image = Image.new("RGB", (600, 200), "white")
    draw = ImageDraw.Draw(image)
    draw.text((20, 20), "scanned timetable", fill="black")
    img_path = str(tmp_path / "scan.png")
    image.save(img_path)
    pdf_path = str(tmp_path / "scan.pdf")
    doc = canvas.Canvas(pdf_path)
    doc.drawImage(img_path, 50, 500, width=400, height=130)
    doc.save()

    def fake_ocr(image):
        return [{"text": "MON", "x0": 120, "y0": 0, "x1": 160,
                 "y1": 12, "conf": 90.0},
                {"text": "TUE", "x0": 220, "y0": 0, "x1": 260,
                 "y1": 12, "conf": 90.0},
                {"text": "09:00-10:00", "x0": 0, "y0": 20, "x1": 80,
                 "y1": 32, "conf": 90.0},
                {"text": "JAVA", "x0": 120, "y0": 20, "x1": 160,
                 "y1": 32, "conf": 90.0},
                {"text": "OS", "x0": 220, "y0": 20, "x1": 260,
                 "y1": 32, "conf": 90.0}]

    monkeypatch.setattr(image_importer, "ocr_words", fake_ocr)
    extraction = TimetableImporter.import_file(pdf_path)
    assert "pdf-ocr" in extraction.method
    assert extraction.lectures == 2


def test_corrupt_pdf_and_image(tmp_path):
    bad_pdf = tmp_path / "bad.pdf"
    bad_pdf.write_bytes(b"%PDF-1.4 not really a pdf")
    with pytest.raises(TimetableImportError):
        TimetableImporter.import_file(str(bad_pdf))
    bad_png = tmp_path / "bad.png"
    bad_png.write_bytes(b"\x89PNG not an image")
    with pytest.raises(TimetableImportError):
        TimetableImporter.import_file(str(bad_png))


# ---- routing, errors, offline ----

def test_unsupported_extension():
    with pytest.raises(TimetableImportError, match="not supported"):
        TimetableImporter.import_file("notes.doc")
    with pytest.raises(TimetableImportError, match="not supported"):
        TimetableImporter.import_file("notes.txt")


def test_mixed_format_import_and_train(tmp_path, monkeypatch):
    import socket
    import urllib.request

    def _fail(*args, **kwargs):
        raise OSError("network blocked")

    monkeypatch.setattr(socket, "socket", _fail)
    monkeypatch.setattr(socket, "create_connection", _fail)
    monkeypatch.setattr(urllib.request, "urlopen", _fail)
    csv_path = write_csv(str(tmp_path / "a.csv"), ROWS_A)
    xlsx_path = write_xlsx(str(tmp_path / "b.xlsx"), GRID_B)
    analyzed = TimetableImporter.analyze_files([csv_path, xlsx_path])
    assert not analyzed["errors"]
    assert sum(r["lectures"] for r in analyzed["reports"]) >= 10
    dicts, approved, rejected = TimetableImporter.approved_dicts(
        analyzed["extractions"])
    assert len(approved) >= 10
    agent = TimetableAgent(data_dir=tmp_path / "model")
    report = agent.train_records(dicts, source_label="mixed-fixtures")
    assert report["reload_check"] is True
    assert agent.model_status()["trained"] is True


def test_imported_records_reach_generation(tmp_path, monkeypatch):
    import socket
    import urllib.request

    def _fail(*args, **kwargs):
        raise OSError("network blocked")

    monkeypatch.setattr(socket, "socket", _fail)
    monkeypatch.setattr(socket, "create_connection", _fail)
    monkeypatch.setattr(urllib.request, "urlopen", _fail)
    monkeypatch.setattr(OllamaClient, "generate", _fail)
    monkeypatch.setattr(OllamaClient, "ensure_ready", _fail)
    csv_path = write_csv(str(tmp_path / "a.csv"), ROWS_A)
    dicts, _, _ = TimetableImporter.approved_dicts(
        [TimetableImporter.import_file(csv_path)])
    agent = TimetableAgent(data_dir=tmp_path / "model")
    agent.train_records(dicts, source_label="a.csv")
    s = make_session()
    sem = seed(s)
    result = agent.generate_dry_run(s, sem.id, "fill", planner="trained")
    assert result.accepted
    assert ConflictService.detect_all_conflicts(s) == []
    s.close()


def test_baseline_source_labels_are_basenames(tmp_path):
    deep = tmp_path / "Users" / "someone" / "college"
    deep.mkdir(parents=True)
    csv_path = write_csv(str(deep / "Timetable2026.csv"), ROWS_A)
    analyzed = TimetableImporter.analyze_files([csv_path])
    label = TimetableImporter.baseline_source_label(analyzed["extractions"])
    assert label == "Timetable2026.csv"
    assert "someone" not in label


def test_production_gate_still_holds(monkeypatch):
    import build
    monkeypatch.setenv("REQUIRE_BASELINE", "1")
    # Repo ships the sample baseline: production gate must refuse it.
    with pytest.raises(SystemExit):
        build.verify_baseline_package()
    monkeypatch.delenv("REQUIRE_BASELINE")


def test_spec_and_installer_runtime():
    root = Path(__file__).resolve().parents[1]
    spec = (root / "CollegeTimetable.spec").read_text(encoding="utf-8")
    for needle in ("timetable_import", "PIL", "pytesseract", "pymupdf",
                   "xlrd", "cv2", "import_preview_dialog"):
        assert needle in spec, needle
    iss = (root / "installer" / "CollegeTimetableSetup.iss").read_text(
        encoding="utf-8")
    sources = [line for line in iss.splitlines()
               if line.strip().lower().startswith("source:")]
    bundled = " ".join(sources).lower()
    for forbidden in ("timetable.db", "training_rows.jsonl",
                      "timetable_learning_profile", "feedback"):
        assert forbidden not in bundled


def test_pc_isolation_with_imports(tmp_path):
    a_csv = write_csv(str(tmp_path / "a.csv"), ROWS_A)
    b_csv = write_grid_csv(str(tmp_path / "b.csv"), GRID_B)
    a_dicts, _, _ = TimetableImporter.approved_dicts(
        [TimetableImporter.import_file(a_csv)])
    b_dicts, _, _ = TimetableImporter.approved_dicts(
        [TimetableImporter.import_file(b_csv)])
    TimetableAgent(data_dir=tmp_path / "a").train_records(a_dicts, "a.csv")
    TimetableAgent(data_dir=tmp_path / "b").train_records(b_dicts, "b.csv")
    assert TimetableAgent(data_dir=tmp_path / "a").model_status()["stored_rows"] == 5
    assert TimetableAgent(data_dir=tmp_path / "b").model_status()["stored_rows"] == 7
